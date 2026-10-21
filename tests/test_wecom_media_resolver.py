from __future__ import annotations

import base64
import json
import sqlite3
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from worker.pipeline.exporter import export_day
from worker.reply_runtime.message_source import LocalWeComMessageSource
from worker.wecom import media
from worker.wecom import cache_messages


ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _account_root(directory: str) -> Path:
    root = Path(directory) / "account"
    (root / "Data").mkdir(parents=True)
    (root / "Cache" / "Image" / "2026-10").mkdir(parents=True)
    return root


def _write_config(root: Path) -> tuple[Path, Path]:
    keys_path = root.parent / "keys.json"
    keys_path.write_text(
        json.dumps({"file.db": {"enc_key": "11" * 16}}), encoding="utf-8"
    )
    config_path = root.parent / "config.json"
    config_path.write_text(
        json.dumps(
            {
                "wxwork_db_dir": str(root / "Data"),
                "wxwork_keys_file": str(keys_path),
            }
        ),
        encoding="utf-8",
    )
    return config_path, keys_path


def _plaintext_connection(config, db_name, include_wal=True):
    del include_wal
    connection = sqlite3.connect(Path(config["wxwork_db_dir"]) / db_name)
    connection.row_factory = sqlite3.Row
    try:
        yield connection
    finally:
        connection.close()


@contextmanager
def _plaintext_context(config, db_name, include_wal=True):
    connection = sqlite3.connect(Path(config["wxwork_db_dir"]) / db_name)
    connection.row_factory = sqlite3.Row
    try:
        yield connection
    finally:
        connection.close()


def _file_db_with_image(
    root: Path,
    *,
    conversation_id: str = "room",
    message_id: int = 7,
    md5: str = "a" * 32,
    server_id: str = "server-token",
    name: str = "capture.png",
    cache_rel: str = r"2026-10\capture.png",
) -> Path:
    file_db = root / "Data" / "file.db"
    connection = sqlite3.connect(file_db)
    try:
        connection.execute(
            """CREATE TABLE file_table4 (
                   origin INTEGER, message_id INTEGER, file_index INTEGER,
                   message_type INTEGER, extension_type INTEGER, server_id TEXT,
                   server_type INTEGER, name TEXT, size INTEGER,
                   receive_time INTEGER, sender_id INTEGER, conversation_id TEXT,
                   collection_id INTEGER, info_extension BLOB, url TEXT,
                   flags INTEGER, md5 TEXT, last_modify_time INTEGER)"""
        )
        connection.execute(
            "INSERT INTO file_table4 VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                0, message_id, 0, 4, 0, server_id, 0, name, 68,
                1_000, 42, conversation_id, 0, b"", "", 0, md5, 1_000,
            ),
        )
        connection.commit()
    finally:
        connection.close()
    mapping_dir = root / "CacheMapping"
    mapping_dir.mkdir(parents=True, exist_ok=True)
    mapping_db = mapping_dir / "mapping.db"
    connection = sqlite3.connect(mapping_db)
    try:
        connection.execute(
            """CREATE TABLE mapping (
                   type INTEGER DEFAULT 0 NOT NULL, key TEXT DEFAULT '' NOT NULL,
                   file_name TEXT DEFAULT '', last_modify_time INTEGER DEFAULT 0 NOT NULL,
                   file_md5 INTEGER DEFAULT 0 NOT NULL, PRIMARY KEY(type,key))"""
        )
        connection.execute(
            "INSERT INTO mapping VALUES (?,?,?,?,?)",
            (2, server_id, cache_rel, 1_000, 0),
        )
        connection.commit()
    finally:
        connection.close()
    return file_db


def _config(root: Path) -> dict:
    return {
        "wxwork_db_dir": str(root / "Data"),
        "wxwork_keys_file": str(root.parent / "keys.json"),
    }


class MediaResolverTests(unittest.TestCase):
    def test_local_extra_path_is_validated_and_described(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            image = root / "Cache" / "Image" / "2026-10" / "capture.png"
            image.write_bytes(ONE_PIXEL_PNG)
            resolver = media.MediaResolver({"wxwork_db_dir": str(root / "Data")}, process_running=True)
            request = media.MediaRequest(
                message_id=7,
                conversation_id="room",
                content_type=4,
                local_extra=str(image).encode("utf-8"),
                expected_count=1,
            )
            descriptor = resolver.resolve_one(request)[0]
        self.assertTrue(descriptor.available)
        self.assertEqual(descriptor.status, "ready")
        self.assertEqual(descriptor.path, str(image))
        self.assertEqual(descriptor.source, "local_extra_content")
        self.assertEqual(descriptor.mime, "image/png")
        self.assertGreater(descriptor.size, 0)

    def test_same_message_descriptor_matches_reply_and_daily_chains(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            image = root / "Cache" / "Image" / "2026-10" / "capture.png"
            image.write_bytes(ONE_PIXEL_PNG)
            _file_db_with_image(root)
            config_path, _ = _write_config(root)
            config = _config(root)
            raw = {
                "send_time": 1_000,
                "sequence": 1,
                "message_id": 7,
                "server_id": 9,
                "content_type": 2,
                "content_raw": b"",
                "extra_content_raw": b"",
                "local_extra_content_raw": str(image).encode("utf-8"),
            }
            reply_message = {
                "messageId": "7",
                "contentType": "image",
                "sendTime": 1_000,
                "images": [
                    {
                        "filename": "",
                        "mimeType": "image/jpeg",
                        "errorCode": "IMAGE_RESOLUTION_PENDING",
                    }
                ],
            }
            source = LocalWeComMessageSource(config_path)
            try:
                with patch("worker.wecom.local_db.decrypted_connection", _plaintext_context):
                    reply_images = source.refresh_images({"groupId": "room"}, [reply_message])[0]["images"]
            finally:
                source.close()
            resolver = media.MediaResolver(config, process_running=True)
            with patch("worker.wecom.local_db.decrypted_connection", _plaintext_context):
                daily = resolver.resolve_one(
                    media.MediaRequest(
                        message_id=7,
                        conversation_id="room",
                        content_type=2,
                        local_extra=raw["local_extra_content_raw"],
                        expected_count=1,
                    )
                )[0]
        self.assertTrue(daily.available)
        self.assertEqual(reply_images[0]["localPath"], daily.path)
        self.assertEqual(reply_images[0]["filename"], daily.filename)
        self.assertEqual(reply_images[0]["mimeType"], daily.mime)

    def test_md5_reference_resolves_forwarded_image_across_conversations(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            image = root / "Cache" / "Image" / "2026-10" / "capture.png"
            image.write_bytes(ONE_PIXEL_PNG)
            _file_db_with_image(root, conversation_id="source-room", md5="f" * 32)
            resolver = media.MediaResolver(_config(root), process_running=True)
            with patch("worker.wecom.local_db.decrypted_connection", _plaintext_context):
                descriptors = resolver.resolve_one(
                    media.MediaRequest(
                        message_id=7,
                        conversation_id="target-room",
                        content_type=123,
                        content_raw=b"\x52\x20" + b"f" * 32,
                        expected_count=1,
                        md5_refs=("f" * 32,),
                    )
                )
        self.assertTrue(descriptors[0].available)
        self.assertEqual(descriptors[0].path, str(image))

    def test_partial_empty_and_oversized_files_get_distinct_reasons(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            cache = root / "Cache" / "Image" / "2026-10"
            truncated = cache / "truncated.jpg"
            truncated.write_bytes(b"\xff\xd8\xff\xe0" + b"x" * 40)
            empty = cache / "empty.png"
            empty.write_bytes(b"")
            oversized = cache / "huge.png"
            with oversized.open("wb") as handle:
                handle.truncate(media.MAX_IMAGE_BYTES + 1)
            resolver = media.MediaResolver({"wxwork_db_dir": str(root / "Data")}, process_running=True)
            for path, expected in (
                (truncated, media.REASON_PARTIAL),
                (empty, media.REASON_PARTIAL),
                (oversized, media.REASON_TOO_LARGE),
            ):
                descriptor = media.inspect_image_file(path)
                self.assertFalse(descriptor.available)
                self.assertEqual(descriptor.reason, expected)

    def test_unsupported_format_and_outside_cache_path_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            outside = Path(directory) / "outside.png"
            outside.write_bytes(ONE_PIXEL_PNG)
            descriptor = media.inspect_image_file(outside)
            self.assertTrue(descriptor.available)  # direct inspection trusts bytes
            config = {"wxwork_db_dir": str(root / "Data")}
            self.assertEqual(media.extract_local_image_paths(config, str(outside).encode()), [])
            icon = root / "Cache" / "Image" / "2026-10" / "icon.svg"
            icon.write_bytes(b"<svg/>")
            self.assertEqual(
                media.inspect_image_file(icon).reason, media.REASON_UNSUPPORTED
            )

    def test_preconditions_return_actionable_reasons(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            (root / "Data" / "session.db").write_bytes(b"x")
            no_dir = media.MediaResolver({})
            self.assertEqual(no_dir.precondition().reason, media.REASON_NOT_FOUND)
            self.assertTrue(no_dir.precondition().action)
            config_path, keys_path = _write_config(root)
            config = _config(root)
            config["wxwork_keys_file"] = str(Path(directory) / "missing-keys.json")
            resolver = media.MediaResolver(config)
            self.assertEqual(resolver.precondition().reason, media.REASON_KEYS_MISSING)
            config["wxwork_keys_file"] = str(keys_path)
            running = media.MediaResolver(config, process_running=False)
            self.assertEqual(running.wait_blocked_reason(), media.REASON_WECOM_NOT_RUNNING)
            self.assertTrue(config_path.exists())

    def test_missing_cache_does_not_poll_when_wecom_is_not_running(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            (root / "Data" / "session.db").write_bytes(b"x")
            _write_config(root)

            def fail_sleep(_seconds):
                raise AssertionError("must not sleep when the precondition blocks")

            resolver = media.MediaResolver(_config(root), process_running=False)
            descriptor = resolver.resolve_batch(
                [media.MediaRequest(message_id=7, conversation_id="room", content_type=4, expected_count=1)],
                wait_seconds=5,
                sleep=fail_sleep,
            )[7][0]
        self.assertFalse(descriptor.available)
        self.assertEqual(descriptor.reason, media.REASON_WECOM_NOT_RUNNING)

    def test_wait_recovers_an_image_written_during_the_budget(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            (root / "Data" / "session.db").write_bytes(b"x")
            image = root / "Cache" / "Image" / "2026-10" / "late.png"
            _write_config(root)

            def sleep(seconds):
                del seconds
                image.write_bytes(ONE_PIXEL_PNG)

            resolver = media.MediaResolver(_config(root), process_running=True)
            descriptors = resolver.resolve_batch(
                [
                    media.MediaRequest(
                        message_id=7,
                        conversation_id="room",
                        content_type=4,
                        local_extra=str(image).encode("utf-8"),
                        expected_count=1,
                        send_time=int(__import__("time").time()),
                    )
                ],
                wait_seconds=5,
                sleep=sleep,
            )[7]
        self.assertTrue(descriptors[0].available)
        self.assertEqual(descriptors[0].path, str(image))

    def test_model_limits_report_too_many_and_too_large(self):
        descriptors = [
            media.MediaDescriptor.ready(Path("a.png"), size=1) for _ in range(9)
        ]
        self.assertEqual(media.model_limit_reason(descriptors), media.REASON_TOO_MANY)
        big = [media.MediaDescriptor.ready(Path("a.png"), size=media.MAX_IMAGE_TOTAL_BYTES + 1)]
        self.assertEqual(media.model_limit_reason(big), media.REASON_TOO_LARGE)


class RescanAndExportTests(unittest.TestCase):
    def _day_dir(self, directory: str, *, with_image: bool) -> tuple[Path, Path]:
        day_dir = Path(directory) / "work" / "2026-10-21"
        day_dir.mkdir(parents=True)
        record = {
            "date": "2026-10-21",
            "message_time": "2026-10-21 10:00:00",
            "sender": "Alice",
            "sender_id": 42,
            "raw_text": "[图片] 帮忙看一下",
            "dedupe_key": "room:7:9:1000",
            "conversation_id": "room",
            "conversation_name": "测试群",
            "message_id": 7,
            "server_id": 9,
            "sequence": 1,
            "content_type": 4,
            "type": "图片",
            "send_time": 1_000,
            "images": [],
            "image_paths": [],
            "image_count_visible": 1,
            "image_status": "missing_original_images",
            "missing_image_names": [],
            "missing_images": [{"reason": media.REASON_NOT_CACHED, "filename": "", "available": False}],
            "missing_image_reasons": [media.REASON_NOT_CACHED],
            "files": [],
            "is_video": False,
            "source_row_index": 1,
            "source_table": "message_table",
        }
        (day_dir / "raw_messages.jsonl").write_text(
            json.dumps(record, ensure_ascii=False) + "\n", encoding="utf-8"
        )
        bulk = day_dir / "raw_attachments" / "_bulk_hd_cache"
        bulk.mkdir(parents=True)
        (bulk / "hd_cache_manifest.json").write_text(
            json.dumps({"date": "2026-10-21", "records": []}), encoding="utf-8"
        )
        image = Path(directory) / "account" / "Cache" / "Image" / "2026-10" / "late.png"
        if with_image:
            image.parent.mkdir(parents=True, exist_ok=True)
            image.write_bytes(ONE_PIXEL_PNG)
        return day_dir, image

    def test_rescan_recovers_late_image_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            _write_config(root)
            config_path = root.parent / "config.json"
            day_dir, image = self._day_dir(directory, with_image=True)
            raw = {
                "message_id": 7,
                "content_type": 4,
                "local_extra_content_raw": str(image).encode("utf-8"),
                "content_raw": b"",
                "extra_content_raw": b"",
                "send_time": 1_000,
            }
            with patch.object(cache_messages, "_rescan_raw_messages", return_value={7: raw}):
                first = cache_messages.rescan_images(day_dir, config_path=config_path)
                second = cache_messages.rescan_images(day_dir, config_path=config_path)
            records = [json.loads(line) for line in (day_dir / "raw_messages.jsonl").read_text(encoding="utf-8").splitlines() if line]
            manifest = json.loads(
                (day_dir / "raw_attachments" / "_bulk_hd_cache" / "hd_cache_manifest.json").read_text(encoding="utf-8")
            )
        self.assertTrue(first["changed"])
        self.assertEqual(first["availableImages"], 1)
        self.assertFalse(second["changed"])
        self.assertEqual(len(records[0]["images"]), 1)
        self.assertEqual(records[0]["image_status"], "ready")
        self.assertEqual(len([item for item in manifest["records"] if item.get("source_message_id") == 7]), 1)

    def test_missing_images_do_not_block_exports_and_are_annotated(self):
        with tempfile.TemporaryDirectory() as directory:
            day_dir, _ = self._day_dir(directory, with_image=False)
            outputs = export_day(day_dir, "2026-10-21", "测试群")
            markdown = Path(outputs["markdown"]).read_text(encoding="utf-8")
            from openpyxl import load_workbook

            workbook = load_workbook(outputs["xlsx"], read_only=True, data_only=True)
            info_rows = {
                str(row[0]): str(row[1])
                for row in workbook["导出说明"].iter_rows(min_row=1, values_only=True)
                if row and row[0]
            }
            chat_headers = [
                cell.value for cell in next(workbook["聊天记录"].iter_rows(min_row=1, max_row=1))
            ]
            workbook.close()
        self.assertIn("图片缺失", markdown)
        self.assertIn("图片尚未写入本机缓存", markdown)
        self.assertIn("图片缺失", info_rows)
        self.assertNotEqual(info_rows["图片缺失"], "无")
        self.assertIn("图片状态", chat_headers)
        self.assertIn("缺失原因", chat_headers)

    def test_smart_sheet_preview_reports_missing_reasons_without_blocking(self):
        from worker.pipeline.smart_sheet import summarize_image_integrity

        summary = summarize_image_integrity(
            [{"key": "issue-a", "expected_image_count": 1, "image_refs": [], "image_status": "missing_original_images"}],
            {},
            mapped=True,
            messages=[
                {
                    "image_status": "missing_original_images",
                    "missing_images": [{"reason": media.REASON_NOT_CACHED, "filename": "a.png"}],
                }
            ],
        )
        self.assertFalse(summary["complete"])
        self.assertEqual(summary["missing"], 1)
        self.assertEqual(summary["reasons"], {media.REASON_NOT_CACHED: 1})
        self.assertEqual(summary["reason_labels"][media.REASON_NOT_CACHED], "图片尚未写入本机缓存")

    def test_cache_messages_main_uses_local_extra_content_for_daily_images(self):
        with tempfile.TemporaryDirectory() as directory:
            root = _account_root(directory)
            (root / "Data" / "session.db").write_bytes(b"x")
            config_path, _ = _write_config(root)
            image = root / "Cache" / "Image" / "2026-10" / "daily.png"
            image.write_bytes(ONE_PIXEL_PNG)
            workspace = Path(directory) / "workspace"
            raw = {
                "source_table": "message_table",
                "message_id": 101,
                "server_id": 201,
                "sequence": 1,
                "sender_id": 301,
                "conversation_id": "room",
                "content_type": 4,
                "send_time": int(__import__("time").time()),
                "flag": 0,
                "content_raw": b"",
                "extra_content_raw": b"",
                "local_extra_content_raw": str(image).encode("utf-8"),
            }
            import contextlib
            import io
            import sys

            stdout = io.StringIO()
            argv = [
                "cache_messages.py",
                "--config",
                str(config_path),
                "--workspace",
                str(workspace),
                "--date",
                "2026-10-21",
                "--conversation-id",
                "room",
            ]
            with (
                patch.object(sys, "argv", argv),
                patch.object(
                    cache_messages,
                    "get_conversation_state",
                    return_value={"display_name": "测试群", "last_message_time": raw["send_time"]},
                ),
                patch.object(cache_messages, "read_messages", return_value=[raw]),
                patch.object(cache_messages, "load_user_map", return_value={301: "Alice"}),
                patch.object(cache_messages, "load_member_names", return_value={}),
                contextlib.redirect_stdout(stdout),
            ):
                exit_code = cache_messages.main()
            payload = json.loads(stdout.getvalue())
            records = [
                json.loads(line)
                for line in Path(payload["raw_messages"]).read_text(encoding="utf-8").splitlines()
                if line
            ]
            manifest = json.loads(Path(payload["image_manifest"]).read_text(encoding="utf-8"))
            self.assertEqual(exit_code, 0)
            self.assertEqual(records[0]["image_status"], "ready")
            self.assertEqual(len(records[0]["images"]), 1)
            self.assertEqual(records[0]["image_descriptors"][0]["source"], "local_extra_content")
            self.assertTrue(Path(records[0]["images"][0]["local_path"]).is_file())
            self.assertEqual(len(manifest["records"]), 1)

    def test_build_media_requests_extracts_refs_and_hints(self):
        records = [
            {
                "message_id": 7,
                "conversation_id": "room",
                "content_type": 123,
                "image_count_visible": 1,
                "raw_text": "截图 see capture.png",
                "missing_image_names": [],
                "send_time": 1_000,
            }
        ]
        raw = {
            "content_type": 123,
            "content_raw": b"\x52\x20" + b"a" * 32,
            "extra_content_raw": b"",
            "local_extra_content_raw": b"",
        }
        requests = cache_messages.build_media_requests(records, {7: raw})
        self.assertEqual(len(requests), 1)
        self.assertEqual(requests[0].md5_refs, ("a" * 32,))
        self.assertIn("capture.png", requests[0].filename_hints)
        self.assertEqual(requests[0].expected_count, 1)


if __name__ == "__main__":
    unittest.main()
